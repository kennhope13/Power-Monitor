import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

md_path = "/home/admin-/Desktop/Power-Monitor/HO_SO_KY_THUAT.md"
docx_path = "/home/admin-/Desktop/Power-Monitor/HO_SO_KY_THUAT.docx"

def set_cell_background(cell, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def add_paragraph_with_spacing(doc, text="", style=None, space_before=0, space_after=6, line_spacing=1.15):
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    return p

def convert_md_to_docx():
    if not os.path.exists(md_path):
        print(f"Error: {md_path} not found.")
        return

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page Setup - Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Đặt font chữ mặc định cho tài liệu (Arial)
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(40, 40, 40)

    in_code_block = False
    code_text = []
    in_table = False
    table_rows = []
    
    first_h2 = True

    i = 0
    while i < len(lines):
        line = lines[i].strip('\r\n')
        stripped = line.strip()

        # Image detection: ![caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            caption = img_match.group(1)
            img_path = img_match.group(2)
            
            # Resolve relative path to absolute
            abs_img_path = img_path
            if not os.path.isabs(img_path):
                abs_img_path = os.path.join("/home/admin-/Desktop/Power-Monitor", img_path)

            if os.path.exists(abs_img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                
                # Add picture
                run_pic = p.add_run()
                run_pic.add_picture(abs_img_path, width=Inches(6.0))
                
                # Caption paragraph
                p_cap = add_paragraph_with_spacing(doc, space_before=2, space_after=12)
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run_cap = p_cap.add_run(f"Hình: {caption}")
                run_cap.font.name = 'Arial'
                run_cap.font.size = Pt(9.5)
                run_cap.italic = True
                run_cap.font.color.rgb = RGBColor(120, 120, 120)
            else:
                print("Warning: image path not found:", abs_img_path)
            i += 1
            continue

        # Code block detection
        if stripped.startswith('```'):
            if in_code_block:
                in_code_block = False
                p = add_paragraph_with_spacing(doc, style='Normal', space_before=6, space_after=6)
                p.paragraph_format.left_indent = Inches(0.4)
                
                run = p.add_run('\n'.join(code_text))
                run.font.name = 'Courier New'
                run.font.size = Pt(9.0)
                run.font.color.rgb = RGBColor(60, 60, 60)
                code_text = []
            else:
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_text.append(line)
            i += 1
            continue
            
        # Table detection
        if stripped.startswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            parsed_rows = []
            for r in table_rows:
                cells = [c.strip() for c in r.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                    continue
                parsed_rows.append(cells)
                
            if parsed_rows:
                cols_count = len(parsed_rows[0])
                table = doc.add_table(rows=len(parsed_rows), cols=cols_count)
                table.style = 'Table Grid'
                table.autofit = True
                
                for row_idx, row_cells in enumerate(parsed_rows):
                    for col_idx, cell_text in enumerate(row_cells):
                        if col_idx < len(table.rows[row_idx].cells):
                            cell = table.rows[row_idx].cells[col_idx]
                            cell.text = cell_text
                            if row_idx == 0:
                                set_cell_background(cell, "E2E8F0") # Light gray header
                                for p in cell.paragraphs:
                                    p.paragraph_format.space_before = Pt(4)
                                    p.paragraph_format.space_after = Pt(4)
                                    for run in p.runs:
                                        run.bold = True
                                        run.font.name = 'Arial'
                                        run.font.size = Pt(9.5)
                            else:
                                for p in cell.paragraphs:
                                    p.paragraph_format.space_before = Pt(3)
                                    p.paragraph_format.space_after = Pt(3)
                                    for run in p.runs:
                                        run.font.name = 'Arial'
                                        run.font.size = Pt(9.0)
                # Add spacing after table
                doc.add_paragraph()
                
            in_table = False
            table_rows = []

        # Headers
        if stripped.startswith('# '):
            title = stripped[2:]
            p = add_paragraph_with_spacing(doc, space_before=24, space_after=12)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(title)
            run.font.name = 'Arial'
            run.font.size = Pt(18)
            run.bold = True
            run.font.color.rgb = RGBColor(16, 44, 87)
        elif stripped.startswith('## '):
            h = stripped[3:]
            if not first_h2:
                doc.add_page_break()
            else:
                first_h2 = False
                
            p = add_paragraph_with_spacing(doc, space_before=18, space_after=6)
            run = p.add_run(h)
            run.font.name = 'Arial'
            run.font.size = Pt(14)
            run.bold = True
            run.font.color.rgb = RGBColor(53, 114, 239)
        elif stripped.startswith('### '):
            h = stripped[4:]
            p = add_paragraph_with_spacing(doc, space_before=12, space_after=4)
            run = p.add_run(h)
            run.font.name = 'Arial'
            run.font.size = Pt(12)
            run.bold = True
            run.font.color.rgb = RGBColor(50, 50, 50)
        elif stripped.startswith('#### '):
            h = stripped[5:]
            p = add_paragraph_with_spacing(doc, space_before=8, space_after=2)
            run = p.add_run(h)
            run.font.name = 'Arial'
            run.font.size = Pt(11)
            run.bold = True
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 100, 100)
            
        # Bullet lists
        elif stripped.startswith('* ') or stripped.startswith('- '):
            indent_level = len(line) - len(line.lstrip())
            text = stripped[2:]
            p = add_paragraph_with_spacing(doc, style='List Bullet', space_before=2, space_after=2)
            p.paragraph_format.left_indent = Inches(0.25 + (indent_level * 0.15))
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(11)
                
        # Numbered lists
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            p = add_paragraph_with_spacing(doc, style='List Number', space_before=2, space_after=2)
            parts = re.split(r'(\*\*.*?\*\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    p.add_run(part)
            for run in p.runs:
                run.font.name = 'Arial'
                run.font.size = Pt(11)
                
        # Blockquotes/Alerts
        elif stripped.startswith('>'):
            clean_alert = re.sub(r'^>\s*\[\![A-Z]+\s*\]\s*', '', stripped)
            clean_alert = re.sub(r'^>\s*', '', clean_alert)
            if clean_alert:
                p = add_paragraph_with_spacing(doc, space_before=4, space_after=4)
                p.paragraph_format.left_indent = Inches(0.4)
                run = p.add_run(clean_alert)
                run.font.name = 'Arial'
                run.font.size = Pt(10.5)
                run.italic = True
                run.font.color.rgb = RGBColor(128, 0, 0)
                
        # Horizontal rules
        elif stripped == '---':
            p = add_paragraph_with_spacing(doc, space_before=12, space_after=12)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("____________________________________________________")
            run.font.color.rgb = RGBColor(200, 200, 200)
            
        # Normal paragraph
        elif stripped:
            p = add_paragraph_with_spacing(doc, space_before=4, space_after=6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            parts = re.split(r'(\*\*.*?\*\*|`.*?`)', stripped)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(180, 50, 50)
                else:
                    p.add_run(part)
            for run in p.runs:
                if not run.font.name:
                    run.font.name = 'Arial'
                run.font.size = Pt(11)
                
        i += 1

    doc.save(docx_path)
    print("DOCX generated successfully at:", docx_path)

if __name__ == '__main__':
    convert_md_to_docx()
